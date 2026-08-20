from pathlib import Path

from docx import Document as DocxDocument
from openpyxl import Workbook
from pypdf import PdfWriter
from reportlab.pdfgen.canvas import Canvas

from app.parsers import get_parser
from app.parsers.pdf import PdfParser, score_text_quality


def test_txt_and_markdown_parsing(tmp_path: Path) -> None:
    txt = tmp_path / "notes.txt"
    txt.write_text("I2C bus recovery\nCheck pull-up resistors.", encoding="utf-8")
    txt_docs = get_parser(txt).parse(txt)
    assert txt_docs[0].content.startswith("I2C")
    assert txt_docs[0].metadata["file_type"] == "txt"

    md = tmp_path / "guide.md"
    md.write_text("# Setup\n\nConfigure pins.\n\n## Faults\n\nCheck ACK.", encoding="utf-8")
    md_docs = get_parser(md).parse(md)
    assert [doc.metadata["section"] for doc in md_docs] == ["Setup", "Faults"]
    assert "Check ACK" in md_docs[1].content


def test_pdf_text_and_page_metadata(tmp_path: Path) -> None:
    path = tmp_path / "datasheet.pdf"
    canvas = Canvas(str(path))
    canvas.drawString(72, 760, "4.8.4 I2C Interface")
    canvas.drawString(72, 740, "The controller supports standard and fast mode operation.")
    canvas.save()

    docs = PdfParser(min_chars_per_page=10).parse(path)
    assert len(docs) == 1
    assert "I2C Interface" in docs[0].content
    assert docs[0].metadata["page"] == 1
    assert docs[0].metadata["section"] == "4.8.4 I2C Interface"
    assert docs[0].metadata["ocr_required"] is False
    assert docs[0].metadata["extraction_method"] == "pdfium"
    assert docs[0].metadata["text_quality_score"] >= 0.65


def test_scanned_or_empty_pdf_is_explicitly_flagged(tmp_path: Path) -> None:
    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with path.open("wb") as stream:
        writer.write(stream)
    docs = PdfParser(min_chars_per_page=10).parse(path)
    assert docs[0].content == ""
    assert docs[0].metadata["ocr_required"] is True


def test_quality_score_rejects_fragmented_garbled_scripts() -> None:
    readable = "ESP32 有两个 I2C 总线接口，支持标准模式和高速模式。"
    garbled = "ESP32 Ⴕ 2۱I2C१ॖၛႨቔ I2Cଆൔ หྟ ᆦӻѓሙଆൔ"
    assert score_text_quality(readable) >= 0.65
    assert score_text_quality(garbled) < score_text_quality(readable)


def test_docx_preserves_heading_paragraph_and_table(tmp_path: Path) -> None:
    path = tmp_path / "manual.docx"
    document = DocxDocument()
    document.add_heading("Sensor Faults", level=1)
    document.add_paragraph("Verify SDA and SCL voltage levels.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Code"
    table.cell(0, 1).text = "Meaning"
    table.cell(1, 0).text = "E01"
    table.cell(1, 1).text = "No ACK"
    document.save(path)

    docs = get_parser(path).parse(path)
    assert docs[0].metadata["section"] == "Sensor Faults"
    assert docs[0].metadata["block_type"] == "paragraph"
    assert docs[1].metadata["block_type"] == "table"
    assert "E01 | No ACK" in docs[1].content


def test_xlsx_emits_one_structured_document_per_row(tmp_path: Path) -> None:
    path = tmp_path / "faults.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Cases"
    sheet.append(["fault", "cause", "action"])
    sheet.append(["No response", "Missing pull-up", "Add 4.7k resistor"])
    sheet.append(["Timeout", "Wrong address", "Scan bus"])
    workbook.save(path)

    docs = get_parser(path).parse(path)
    assert len(docs) == 2
    assert docs[0].metadata["sheet"] == "Cases"
    assert docs[0].metadata["row"] == 2
    assert docs[0].metadata["record_type"] == "row"
    assert "cause: Missing pull-up" in docs[0].content
