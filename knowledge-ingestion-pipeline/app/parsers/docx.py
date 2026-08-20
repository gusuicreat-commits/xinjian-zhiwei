from __future__ import annotations

from pathlib import Path
from typing import List, Optional

from docx import Document as DocxDocument
from docx.document import Document as DocxDocumentType
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.models import Document
from app.parsers.base import Parser


def _iter_blocks(parent: DocxDocumentType):
    for child in parent.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, parent)
        elif child.tag.endswith("}tbl"):
            yield Table(child, parent)


class DocxParser(Parser):
    def parse(self, path: Path) -> List[Document]:
        doc = DocxDocument(str(path))
        base = self.base_metadata(path, "docx")
        result: List[Document] = []
        section: Optional[str] = None
        paragraph_buffer: List[str] = []

        def flush() -> None:
            if paragraph_buffer:
                metadata = {**base, "block_type": "paragraph"}
                if section:
                    metadata["section"] = section
                result.append(Document("\n\n".join(paragraph_buffer), metadata))
                paragraph_buffer.clear()

        for block in _iter_blocks(doc):
            if isinstance(block, Paragraph):
                text = block.text.strip()
                if not text:
                    flush()
                    continue
                if block.style and block.style.name.lower().startswith("heading"):
                    flush()
                    section = text
                else:
                    paragraph_buffer.append(text)
            else:
                flush()
                rows = [[cell.text.strip() for cell in row.cells] for row in block.rows]
                if not rows:
                    continue
                table_text = "\n".join(" | ".join(row) for row in rows)
                metadata = {**base, "block_type": "table", "row_count": len(rows)}
                if section:
                    metadata["section"] = section
                result.append(Document(table_text, metadata))
        flush()
        return result
