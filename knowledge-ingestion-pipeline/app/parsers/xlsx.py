from __future__ import annotations

from pathlib import Path
from typing import List

from openpyxl import load_workbook

from app.models import Document
from app.parsers.base import Parser


class XlsxParser(Parser):
    def parse(self, path: Path) -> List[Document]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        base = self.base_metadata(path, "xlsx")
        documents: List[Document] = []
        try:
            for sheet in workbook.worksheets:
                rows = list(sheet.iter_rows(values_only=True))
                if not rows:
                    continue
                headers = self._headers(rows[0])
                for row_number, row in enumerate(rows[1:], start=2):
                    values = list(row)
                    if not any(value is not None and str(value).strip() for value in values):
                        continue
                    pairs = [
                        f"{headers[index]}: {value}"
                        for index, value in enumerate(values)
                        if value is not None and str(value).strip()
                    ]
                    documents.append(
                        Document(
                            content="\n".join(pairs),
                            metadata={
                                **base,
                                "sheet": sheet.title,
                                "row": row_number,
                                "columns": headers[: len(values)],
                                "record_type": "row",
                            },
                        )
                    )
        finally:
            workbook.close()
        return documents

    @staticmethod
    def _headers(row: tuple) -> List[str]:
        return [
            str(value).strip() if value is not None else f"column_{i + 1}"
            for i, value in enumerate(row)
        ]
